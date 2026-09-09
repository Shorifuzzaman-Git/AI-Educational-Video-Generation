import os
import re
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.text import MSO_AUTO_SIZE

from PIL import Image

from services.image_search import ImageSearchService
# from services.voice_generate import VoiceGenerator

# from services.image_search import ImageSearchService


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent

GENERATED_DIR = BASE_DIR / "generated"
SLIDES_DIR = GENERATED_DIR / "slides"
IMAGE_DIR = GENERATED_DIR / "images"

SLIDES_DIR.mkdir(parents=True, exist_ok=True)
IMAGE_DIR.mkdir(parents=True, exist_ok=True)


# PowerPoint 16:9
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


# ============================================================
# COLORS
# ============================================================

NAVY = RGBColor(20, 35, 55)
DARK = RGBColor(35, 42, 50)
GRAY = RGBColor(105, 112, 120)
LIGHT_GRAY = RGBColor(238, 241, 244)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(20, 20, 20)
GREEN = RGBColor(43, 130, 95)
BLUE = RGBColor(52, 100, 180)
ORANGE = RGBColor(220, 130, 50)
RED = RGBColor(190, 70, 70)


# ============================================================
# MAIN FUNCTION
# ============================================================

def create_slides(
    presentation_data: dict,
    filename: str = "generated_presentation.pptx"
) -> str:
    """
    Create a PowerPoint presentation from structured LLM JSON.

    Parameters
    ----------
    presentation_data : dict
        Structured presentation data returned by LLMService.

    filename : str
        Output PowerPoint filename.

    Returns
    -------
    str
        Path of generated PPTX file.
    """

    if not presentation_data:
        raise ValueError("presentation_data cannot be empty")

    if "slides" not in presentation_data:
        raise ValueError("presentation_data must contain 'slides'")

    slides_data = presentation_data["slides"]

    if not slides_data:
        raise ValueError("No slides found in presentation_data")

    # --------------------------------------------------------
    # Create presentation
    # --------------------------------------------------------

    prs = Presentation()

    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    # Remove default slide
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[-1].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[-1]

    # --------------------------------------------------------
    # Generate each slide
    # --------------------------------------------------------

    for index, slide_data in enumerate(slides_data):

        layout = slide_data.get(
            "layout",
            "image_left_content_right"
        )

        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        # Background
        set_background(slide, WHITE)

        # Select layout
        if layout == "title_hero":
            create_title_hero_slide(
                slide,
                slide_data,
                index
            )

        elif layout == "image_left_content_right":
            create_image_left_content_right_slide(
                slide,
                slide_data,
                index
            )

        elif layout == "dashboard":
            create_dashboard_slide(
                slide,
                slide_data,
                index
            )

        elif layout == "statistics":
            create_statistics_slide(
                slide,
                slide_data,
                index
            )

        elif layout == "cards":
            create_cards_slide(
                slide,
                slide_data,
                index
            )

        elif layout == "comparison":
            create_comparison_slide(
                slide,
                slide_data,
                index
            )

        elif layout == "timeline":
            create_timeline_slide(
                slide,
                slide_data,
                index
            )

        elif layout == "chart":
            create_chart_slide(
                slide,
                slide_data,
                index
            )

        elif layout == "process":
            create_process_slide(
                slide,
                slide_data,
                index
            )

        elif layout == "conclusion":
            create_conclusion_slide(
                slide,
                slide_data,
                index
            )

        else:
            create_image_left_content_right_slide(
                slide,
                slide_data,
                index
            )

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    output_path = SLIDES_DIR / filename

    prs.save(str(output_path))

    print(f"Presentation created: {output_path}")

    return str(output_path)


# ============================================================
# BACKGROUND
# ============================================================

def set_background(slide, color=WHITE):
    """
    Set slide background color.
    """

    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


# ============================================================
# TEXT HELPERS
# ============================================================

def add_text(
    slide,
    text,
    x,
    y,
    width,
    height,
    font_size=24,
    bold=False,
    color=DARK,
    font_name="Aptos",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.05
):
    """
    Add formatted text box.
    """

    textbox = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height)
    )

    text_frame = textbox.text_frame

    text_frame.clear()

    text_frame.margin_left = Inches(margin)
    text_frame.margin_right = Inches(margin)
    text_frame.margin_top = Inches(margin)
    text_frame.margin_bottom = Inches(margin)

    text_frame.vertical_anchor = valign

    paragraph = text_frame.paragraphs[0]

    paragraph.text = str(text)

    paragraph.alignment = align

    run = paragraph.runs[0]

    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color

    return textbox


def add_title(
    slide,
    title,
    subtitle=None
):
    """
    Add standard presentation title.
    """

    add_text(
        slide,
        title,
        0.65,
        0.35,
        12.0,
        0.65,
        font_size=27,
        bold=True,
        color=NAVY
    )

    if subtitle:

        add_text(
            slide,
            subtitle,
            0.68,
            1.02,
            11.5,
            0.45,
            font_size=12,
            color=GRAY
        )


def add_footer(
    slide,
    slide_number,
    source_names=None
):
    """
    Add footer and slide number.
    """

    # Horizontal line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.65),
        Inches(7.12),
        Inches(12.0),
        Inches(0.01)
    )

    line.fill.solid()
    line.fill.fore_color.rgb = LIGHT_GRAY
    line.line.fill.background()

    # Source
    if source_names:

        if isinstance(source_names, list):
            source_text = "Sources: " + ", ".join(
                str(x) for x in source_names
            )
        else:
            source_text = f"Source: {source_names}"

        add_text(
            slide,
            source_text,
            0.68,
            7.16,
            10.5,
            0.2,
            font_size=7,
            color=GRAY
        )

    # Slide number
    add_text(
        slide,
        str(slide_number + 1),
        12.0,
        7.14,
        0.5,
        0.25,
        font_size=8,
        color=GRAY,
        align=PP_ALIGN.RIGHT
    )


# ============================================================
# IMAGE HELPERS
# ============================================================

def get_image_service():
    """
    Create ImageSearchService instance.
    """

    return ImageSearchService()


def search_image(
    query: str,
    filename: Optional[str] = None
):
    """
    Search and download an image.

    Returns local image path or None.
    """

    if not query:
        return None

    try:

        service = get_image_service()

        result = service.search_and_download(
            query=query,
            filename=filename
        )

        if result:
            return result["path"]

    except Exception as e:

        print(
            f"Image search failed for '{query}': {e}"
        )

    return None


def crop_image_to_box(
    image_path,
    output_path,
    target_width,
    target_height
):
    """
    Crop image to target aspect ratio.
    """

    target_ratio = target_width / target_height

    with Image.open(image_path) as img:

        img = img.convert("RGB")

        width, height = img.size

        current_ratio = width / height

        if current_ratio > target_ratio:

            # Too wide
            new_width = int(height * target_ratio)

            left = (width - new_width) // 2

            img = img.crop(
                (
                    left,
                    0,
                    left + new_width,
                    height
                )
            )

        else:

            # Too tall
            new_height = int(width / target_ratio)

            top = (height - new_height) // 2

            img = img.crop(
                (
                    0,
                    top,
                    width,
                    top + new_height
                )
            )

        img = img.resize(
            (
                int(target_width),
                int(target_height)
            )
        )

        img.save(
            output_path,
            quality=95
        )

    return output_path


def add_image_cover(
    slide,
    image_path,
    x,
    y,
    width,
    height
):
    """
    Add image using crop-to-cover behavior.
    """

    if not image_path:
        return None

    try:

        with Image.open(image_path) as img:
            original_width, original_height = img.size

        target_width = int(width * 120)
        target_height = int(height * 120)

        cropped_name = (
            Path(image_path).stem
            + f"_{target_width}x{target_height}.jpg"
        )

        cropped_path = IMAGE_DIR / cropped_name

        if not cropped_path.exists():

            crop_image_to_box(
                image_path,
                str(cropped_path),
                target_width,
                target_height
            )

        picture = slide.shapes.add_picture(
            str(cropped_path),
            Inches(x),
            Inches(y),
            width=Inches(width),
            height=Inches(height)
        )

        return picture

    except Exception as e:

        print(
            f"Could not add image {image_path}: {e}"
        )

        return None


# ============================================================
# CARD
# ============================================================

def add_card(
    slide,
    x,
    y,
    width,
    height,
    title,
    body,
    image_query=None,
    accent=BLUE
):
    """
    Add a modern card.
    """

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height)
    )

    card.fill.solid()
    card.fill.fore_color.rgb = WHITE

    card.line.color.rgb = LIGHT_GRAY

    # Accent bar
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(0.07),
        Inches(height)
    )

    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = accent
    accent_bar.line.fill.background()

    current_y = y + 0.22

    # Image
    if image_query:

        image_path = search_image(
            image_query,
            filename=safe_filename(
                image_query
            )
        )

        if image_path:

            add_image_cover(
                slide,
                image_path,
                x + 0.18,
                y + 0.18,
                width - 0.36,
                1.15
            )

            current_y = y + 1.48

    # Title
    add_text(
        slide,
        title,
        x + 0.22,
        current_y,
        width - 0.4,
        0.45,
        font_size=15,
        bold=True,
        color=NAVY
    )

    # Body
    add_text(
        slide,
        body,
        x + 0.22,
        current_y + 0.48,
        width - 0.4,
        height - (
            current_y - y
        ) - 0.6,
        font_size=10,
        color=GRAY
    )


# ============================================================
# SLIDE 1 - TITLE + HERO
# ============================================================

def create_title_hero_slide(
    slide,
    data,
    index
):
    """
    Title + large hero image.
    """

    title = data.get(
        "title",
        "AI Generated Presentation"
    )

    subtitle = data.get(
        "subtitle",
        ""
    )

    add_text(
        slide,
        title,
        0.75,
        0.75,
        7.0,
        1.4,
        font_size=34,
        bold=True,
        color=NAVY
    )

    add_text(
        slide,
        subtitle,
        0.78,
        2.15,
        6.2,
        1.0,
        font_size=17,
        color=GRAY
    )

    # Hero image
    hero = data.get(
        "hero_image",
        {}
    )

    query = hero.get(
        "image_query",
        data.get(
            "image_query",
            title
        )
    )

    image_path = search_image(
        query,
        filename=f"slide_{index + 1}_hero.jpg"
    )

    if image_path:

        add_image_cover(
            slide,
            image_path,
            7.15,
            0.55,
            5.5,
            5.85
        )

    # Key takeaway
    takeaway = data.get(
        "key_takeaway",
        ""
    )

    if takeaway:

        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.78),
            Inches(4.2),
            Inches(5.8),
            Inches(1.15)
        )

        box.fill.solid()
        box.fill.fore_color.rgb = LIGHT_GRAY
        box.line.fill.background()

        add_text(
            slide,
            "KEY TAKEAWAY",
            1.0,
            4.42,
            2.0,
            0.3,
            font_size=9,
            bold=True,
            color=BLUE
        )

        add_text(
            slide,
            takeaway,
            1.0,
            4.75,
            5.2,
            0.4,
            font_size=13,
            bold=True,
            color=DARK
        )

    add_footer(
        slide,
        index,
        data.get("source_names")
    )


# ============================================================
# SLIDE - IMAGE LEFT / CONTENT RIGHT
# ============================================================

def create_image_left_content_right_slide(
    slide,
    data,
    index
):
    """
    Large image on left and content on right.
    """

    add_title(
        slide,
        data.get("title", ""),
        data.get("subtitle")
    )

    # Image
    hero = data.get(
        "hero_image",
        {}
    )

    query = hero.get(
        "image_query",
        data.get("image_query")
    )

    image_path = search_image(
        query,
        filename=f"slide_{index + 1}_hero.jpg"
    ) if query else None

    if image_path:

        add_image_cover(
            slide,
            image_path,
            0.65,
            1.65,
            5.25,
            4.9
        )

    # Caption
    caption = hero.get(
        "caption"
    )

    if caption:

        add_text(
            slide,
            caption,
            0.8,
            6.62,
            4.95,
            0.35,
            font_size=8,
            color=GRAY,
            align=PP_ALIGN.CENTER
        )

    # Right side
    stats = data.get(
        "statistics",
        []
    )

    cards = data.get(
        "cards",
        []
    )

    current_y = 1.65

    # Statistics
    for stat in stats[:2]:

        add_stat_box(
            slide,
            6.35,
            current_y,
            2.8,
            1.15,
            stat
        )

        current_y += 1.35

    # Cards
    remaining_cards = cards[:3]

    for card in remaining_cards:

        add_small_content_card(
            slide,
            9.35,
            current_y - 1.35,
            3.25,
            1.25,
            card
        )

        current_y += 1.4

    # Key takeaway
    takeaway = data.get(
        "key_takeaway"
    )

    if takeaway:

        add_text(
            slide,
            takeaway,
            6.35,
            5.95,
            6.2,
            0.7,
            font_size=12,
            bold=True,
            color=NAVY
        )

    add_footer(
        slide,
        index,
        data.get("source_names")
    )


# ============================================================
# STATISTICS BOX
# ============================================================

def add_stat_box(
    slide,
    x,
    y,
    width,
    height,
    stat
):
    """
    Add statistic box.
    """

    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height)
    )

    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT_GRAY
    box.line.fill.background()

    value = stat.get(
        "value",
        ""
    )

    label = stat.get(
        "label",
        ""
    )

    description = stat.get(
        "description",
        ""
    )

    add_text(
        slide,
        value,
        x + 0.2,
        y + 0.12,
        width - 0.4,
        0.42,
        font_size=22,
        bold=True,
        color=BLUE
    )

    add_text(
        slide,
        label,
        x + 0.2,
        y + 0.55,
        width - 0.4,
        0.3,
        font_size=10,
        bold=True,
        color=NAVY
    )

    if description:

        add_text(
            slide,
            description,
            x + 0.2,
            y + 0.82,
            width - 0.4,
            0.25,
            font_size=7,
            color=GRAY
        )


# ============================================================
# SMALL CONTENT CARD
# ============================================================

def add_small_content_card(
    slide,
    x,
    y,
    width,
    height,
    card
):
    """
    Small text card.
    """

    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height)
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = LIGHT_GRAY

    title = card.get(
        "title",
        ""
    )

    body = card.get(
        "body",
        ""
    )

    add_text(
        slide,
        title,
        x + 0.2,
        y + 0.15,
        width - 0.4,
        0.3,
        font_size=11,
        bold=True,
        color=NAVY
    )

    add_text(
        slide,
        body,
        x + 0.2,
        y + 0.48,
        width - 0.4,
        height - 0.58,
        font_size=8,
        color=GRAY
    )


# ============================================================
# DASHBOARD SLIDE
# ============================================================

def create_dashboard_slide(
    slide,
    data,
    index
):
    """
    Dashboard-style infographic.
    """

    add_title(
        slide,
        data.get("title", ""),
        data.get("subtitle")
    )

    stats = data.get(
        "statistics",
        []
    )

    # Top statistics
    stats = stats[:4]

    card_width = 2.8
    start_x = 0.65

    for i, stat in enumerate(stats):

        add_stat_box(
            slide,
            start_x + i * 3.05,
            1.55,
            card_width,
            1.25,
            stat
        )

    # Hero image
    hero = data.get(
        "hero_image",
        {}
    )

    query = hero.get(
        "image_query"
    )

    if query:

        image_path = search_image(
            query,
            filename=f"slide_{index + 1}_dashboard.jpg"
        )

        if image_path:

            add_image_cover(
                slide,
                image_path,
                0.65,
                3.05,
                5.0,
                3.25
            )

    # Cards
    cards = data.get(
        "cards",
        []
    )

    for i, card in enumerate(cards[:3]):

        add_small_content_card(
            slide,
            5.95,
            3.05 + i * 1.1,
            6.6,
            0.9,
            card
        )

    add_footer(
        slide,
        index,
        data.get("source_names")
    )


# ============================================================
# STATISTICS SLIDE
# ============================================================

def create_statistics_slide(
    slide,
    data,
    index
):
    """
    Statistics-focused slide.
    """

    add_title(
        slide,
        data.get("title", ""),
        data.get("subtitle")
    )

    stats = data.get(
        "statistics",
        []
    )

    positions = [
        (0.8, 1.7),
        (4.55, 1.7),
        (8.3, 1.7),
        (0.8, 4.1),
        (4.55, 4.1),
        (8.3, 4.1)
    ]

    for i, stat in enumerate(stats[:6]):

        x, y = positions[i]

        add_large_stat_card(
            slide,
            x,
            y,
            3.2,
            1.8,
            stat
        )

    add_footer(
        slide,
        index,
        data.get("source_names")
    )


def add_large_stat_card(
    slide,
    x,
    y,
    width,
    height,
    stat
):
    """
    Large statistic card.
    """

    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height)
    )

    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT_GRAY
    box.line.fill.background()

    value = stat.get(
        "value",
        ""
    )

    label = stat.get(
        "label",
        ""
    )

    description = stat.get(
        "description",
        ""
    )

    add_text(
        slide,
        value,
        x + 0.25,
        y + 0.25,
        width - 0.5,
        0.65,
        font_size=28,
        bold=True,
        color=BLUE,
        align=PP_ALIGN.CENTER
    )

    add_text(
        slide,
        label,
        x + 0.25,
        y + 0.9,
        width - 0.5,
        0.3,
        font_size=12,
        bold=True,
        color=NAVY,
        align=PP_ALIGN.CENTER
    )

    add_text(
        slide,
        description,
        x + 0.25,
        y + 1.25,
        width - 0.5,
        0.35,
        font_size=8,
        color=GRAY,
        align=PP_ALIGN.CENTER
    )


# ============================================================
# CARDS SLIDE
# ============================================================

def create_cards_slide(
    slide,
    data,
    index
):
    """
    Three or four visual cards.
    """

    add_title(
        slide,
        data.get("title", ""),
        data.get("subtitle")
    )

    cards = data.get(
        "cards",
        []
    )

    count = min(
        len(cards),
        4
    )

    if count == 0:
        add_footer(
            slide,
            index,
            data.get("source_names")
        )
        return

    width = 2.9

    for i, card in enumerate(cards[:4]):

        x = 0.65 + i * 3.1

        add_card(
            slide,
            x,
            1.65,
            width,
            4.9,
            card.get("title", ""),
            card.get("body", ""),
            card.get("image_query"),
            [BLUE, GREEN, ORANGE, RED][i % 4]
        )

    add_footer(
        slide,
        index,
        data.get("source_names")
    )


# ============================================================
# COMPARISON SLIDE
# ============================================================

def create_comparison_slide(
    slide,
    data,
    index
):
    """
    Comparison between two subjects.
    """

    add_title(
        slide,
        data.get("title", ""),
        data.get("subtitle")
    )

    cards = data.get(
        "cards",
        []
    )

    if len(cards) >= 2:

        create_comparison_column(
            slide,
            0.75,
            1.65,
            5.8,
            4.9,
            cards[0]
        )

        create_comparison_column(
            slide,
            6.8,
            1.65,
            5.8,
            4.9,
            cards[1]
        )

    else:

        # fallback
        add_text(
            slide,
            "Comparison data unavailable.",
            1.0,
            3.0,
            11,
            1,
            font_size=20,
            color=GRAY,
            align=PP_ALIGN.CENTER
        )

    add_footer(
        slide,
        index,
        data.get("source_names")
    )


def create_comparison_column(
    slide,
    x,
    y,
    width,
    height,
    card
):
    """
    One comparison column.
    """

    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height)
    )

    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT_GRAY
    box.line.fill.background()

    title = card.get(
        "title",
        ""
    )

    body = card.get(
        "body",
        ""
    )

    image_query = card.get(
        "image_query"
    )

    add_text(
        slide,
        title,
        x + 0.3,
        y + 0.25,
        width - 0.6,
        0.5,
        font_size=20,
        bold=True,
        color=NAVY,
        align=PP_ALIGN.CENTER
    )

    if image_query:

        image_path = search_image(
            image_query
        )

        if image_path:

            add_image_cover(
                slide,
                image_path,
                x + 0.4,
                y + 0.9,
                width - 0.8,
                2.0
            )

            body_y = y + 3.15

        else:

            body_y = y + 1.0

    else:

        body_y = y + 1.0

    add_text(
        slide,
        body,
        x + 0.4,
        body_y,
        width - 0.8,
        1.5,
        font_size=11,
        color=DARK
    )


# ============================================================
# TIMELINE SLIDE
# ============================================================

def create_timeline_slide(
    slide,
    data,
    index
):
    """
    Timeline slide.
    """

    add_title(
        slide,
        data.get("title", ""),
        data.get("subtitle")
    )

    cards = data.get(
        "cards",
        []
    )

    if not cards:

        add_footer(
            slide,
            index,
            data.get("source_names")
        )

        return

    # Timeline line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(1.0),
        Inches(3.55),
        Inches(11.2),
        Inches(0.04)
    )

    line.fill.solid()
    line.fill.fore_color.rgb = BLUE
    line.line.fill.background()

    count = min(
        len(cards),
        5
    )

    spacing = 10.8 / max(
        count - 1,
        1
    )

    for i, card in enumerate(cards[:5]):

        x = 1.0 + i * spacing

        # Circle
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x - 0.18),
            Inches(3.35),
            Inches(0.4),
            Inches(0.4)
        )

        circle.fill.solid()
        circle.fill.fore_color.rgb = BLUE
        circle.line.fill.background()

        # Number
        add_text(
            slide,
            str(i + 1),
            x - 0.13,
            3.39,
            0.3,
            0.2,
            font_size=8,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER
        )

        # Title
        add_text(
            slide,
            card.get("title", ""),
            x - 0.75,
            2.25,
            1.7,
            0.7,
            font_size=11,
            bold=True,
            color=NAVY,
            align=PP_ALIGN.CENTER
        )

        # Body
        add_text(
            slide,
            card.get("body", ""),
            x - 0.85,
            4.0,
            1.9,
            1.5,
            font_size=8,
            color=GRAY,
            align=PP_ALIGN.CENTER
        )

    add_footer(
        slide,
        index,
        data.get("source_names")
    )


# ============================================================
# CHART SLIDE
# ============================================================

def create_chart_slide(
    slide,
    data,
    index
):
    """
    Create a simple bar chart using PowerPoint shapes.

    Expected chart format:

    "chart": {
        "title": "...",
        "type": "bar",
        "labels": ["A", "B", "C"],
        "values": [20, 40, 60]
    }
    """

    add_title(
        slide,
        data.get("title", ""),
        data.get("subtitle")
    )

    chart = data.get(
        "chart",
        {}
    )

    labels = chart.get(
        "labels",
        []
    )

    values = chart.get(
        "values",
        []
    )

    chart_title = chart.get(
        "title",
        ""
    )

    if chart_title:

        add_text(
            slide,
            chart_title,
            0.8,
            1.45,
            11.5,
            0.4,
            font_size=14,
            bold=True,
            color=NAVY,
            align=PP_ALIGN.CENTER
        )

    if not labels or not values:

        add_text(
            slide,
            "Chart data unavailable.",
            1,
            3,
            11,
            1,
            font_size=20,
            color=GRAY,
            align=PP_ALIGN.CENTER
        )

        add_footer(
            slide,
            index,
            data.get("source_names")
        )

        return

    # Make sure same length
    count = min(
        len(labels),
        len(values)
    )

    labels = labels[:count]
    values = values[:count]

    max_value = max(
        [float(v) for v in values]
    )

    if max_value <= 0:
        max_value = 1

    chart_x = 1.2
    chart_y = 2.0
    chart_width = 10.8
    chart_height = 3.9

    bar_gap = 0.35

    available_width = (
        chart_width
        - (count - 1) * bar_gap
    )

    bar_width = (
        available_width / count
    )

    for i in range(count):

        value = float(
            values[i]
        )

        bar_height = (
            value / max_value
        ) * 3.2

        x = (
            chart_x
            + i * (bar_width + bar_gap)
        )

        y = (
            chart_y
            + 3.2
            - bar_height
        )

        # Bar
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(bar_width),
            Inches(bar_height)
        )

        bar.fill.solid()
        bar.fill.fore_color.rgb = BLUE
        bar.line.fill.background()

        # Value
        add_text(
            slide,
            format_number(value),
            x,
            y - 0.35,
            bar_width,
            0.3,
            font_size=10,
            bold=True,
            color=NAVY,
            align=PP_ALIGN.CENTER
        )

        # Label
        add_text(
            slide,
            labels[i],
            x,
            chart_y + 3.35,
            bar_width,
            0.5,
            font_size=8,
            color=GRAY,
            align=PP_ALIGN.CENTER
        )

    add_footer(
        slide,
        index,
        data.get("source_names")
    )


# ============================================================
# PROCESS SLIDE
# ============================================================

def create_process_slide(
    slide,
    data,
    index
):
    """
    Process / workflow slide.
    """

    add_title(
        slide,
        data.get("title", ""),
        data.get("subtitle")
    )

    cards = data.get(
        "cards",
        []
    )

    count = min(
        len(cards),
        5
    )

    if count == 0:

        add_footer(
            slide,
            index,
            data.get("source_names")
        )

        return

    box_width = 2.15
    gap = 0.35

    total_width = (
        count * box_width
        + (count - 1) * gap
    )

    start_x = (
        13.333 - total_width
    ) / 2

    for i, card in enumerate(cards[:5]):

        x = start_x + i * (
            box_width + gap
        )

        # Main box
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(2.25),
            Inches(box_width),
            Inches(2.7)
        )

        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = LIGHT_GRAY

        # Number
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x + 0.72),
            Inches(1.72),
            Inches(0.7),
            Inches(0.7)
        )

        circle.fill.solid()
        circle.fill.fore_color.rgb = BLUE
        circle.line.fill.background()

        add_text(
            slide,
            str(i + 1),
            x + 0.72,
            1.91,
            0.7,
            0.2,
            font_size=11,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER
        )

        add_text(
            slide,
            card.get("title", ""),
            x + 0.18,
            2.65,
            box_width - 0.36,
            0.5,
            font_size=14,
            bold=True,
            color=NAVY,
            align=PP_ALIGN.CENTER
        )

        add_text(
            slide,
            card.get("body", ""),
            x + 0.2,
            3.3,
            box_width - 0.4,
            1.3,
            font_size=9,
            color=GRAY,
            align=PP_ALIGN.CENTER
        )

        # Arrow
        if i < count - 1:

            add_text(
                slide,
                "→",
                x + box_width,
                3.25,
                gap,
                0.4,
                font_size=22,
                bold=True,
                color=BLUE,
                align=PP_ALIGN.CENTER
            )

    add_footer(
        slide,
        index,
        data.get("source_names")
    )


# ============================================================
# CONCLUSION SLIDE
# ============================================================

def create_conclusion_slide(
    slide,
    data,
    index
):
    """
    Conclusion / takeaway slide.
    """

    set_background(
        slide,
        LIGHT_GRAY
    )

    title = data.get(
        "title",
        "Conclusion"
    )

    subtitle = data.get(
        "subtitle",
        ""
    )

    add_text(
        slide,
        title,
        1.0,
        1.1,
        11.3,
        0.8,
        font_size=32,
        bold=True,
        color=NAVY,
        align=PP_ALIGN.CENTER
    )

    add_text(
        slide,
        subtitle,
        1.5,
        2.0,
        10.3,
        0.7,
        font_size=16,
        color=GRAY,
        align=PP_ALIGN.CENTER
    )

    takeaway = data.get(
        "key_takeaway",
        ""
    )

    if takeaway:

        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(2.0),
            Inches(3.1),
            Inches(9.3),
            Inches(1.5)
        )

        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.fill.background()

        add_text(
            slide,
            takeaway,
            2.4,
            3.5,
            8.5,
            0.7,
            font_size=18,
            bold=True,
            color=NAVY,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE
        )

    add_footer(
        slide,
        index,
        data.get("source_names")
    )


# ============================================================
# UTILITY FUNCTIONS
# ============================================================

def safe_filename(text):
    """
    Convert arbitrary text into safe filename.
    """

    text = str(text)

    text = re.sub(
        r"[^a-zA-Z0-9_-]+",
        "_",
        text
    )

    text = text.strip(
        "_"
    )

    if not text:
        text = "image"

    return text[:80] + ".jpg"


def format_number(value):
    """
    Format chart number.
    """

    try:

        value = float(value)

        if value.is_integer():

            return f"{int(value):,}"

        return f"{value:,.2f}"

    except Exception:

        return str(value)


# ============================================================
# OPTIONAL: SAVE NARRATION TEXT
# ============================================================

def save_narrations(
    presentation_data: dict,
    filename: str = "narrations.txt"
):
    """
    Save all slide narration into a text file.

    This can later be passed to Edge TTS.
    """

    output_path = GENERATED_DIR / filename

    lines = []

    slides = presentation_data.get(
        "slides",
        []
    )

    for index, slide in enumerate(slides):

        title = slide.get(
            "title",
            f"Slide {index + 1}"
        )

        narration = slide.get(
            "narration",
            ""
        )

        lines.append(
            f"Slide {index + 1}: {title}"
        )

        lines.append(
            narration
        )

        lines.append(
            ""
        )

    output_path.write_text(
        "\n".join(lines),
        encoding="utf-8"
    )

    return str(output_path)


# ============================================================
# TEST
# ============================================================

if __name__ == "__main__":

    test_data = {

        "presentation_title":
            "Agriculture of Bangladesh",

        "slides": [

            {
                "section": "Introduction",

                "title":
                    "Agriculture in Bangladesh",

                "subtitle":
                    "A major pillar of the national economy",

                "layout":
                    "title_hero",

                "narration":
                    "Agriculture remains an important part of Bangladesh's economy and supports millions of people.",

                "hero_image": {
                    "image_query":
                        "Bangladesh rice field farmer agriculture",
                    "caption":
                        "Agriculture remains central to rural livelihoods."
                },

                "statistics": [],

                "cards": [],

                "key_takeaway":
                    "Agriculture supports food security, employment and rural livelihoods.",

                "source_names": [
                    "World Bank",
                    "FAO"
                ]
            },

            {

                "section": "Overview",

                "title":
                    "Agriculture at a Glance",

                "subtitle":
                    "Key indicators",

                "layout":
                    "dashboard",

                "narration":
                    "Bangladesh has a large agricultural sector involving crops, fisheries and livestock.",

                "hero_image": {
                    "image_query":
                        "Bangladesh agriculture rice farming"
                },

                "statistics": [

                    {
                        "label": "Major Crop",
                        "value": "Rice",
                        "description":
                            "Staple food crop"
                    },

                    {
                        "label": "Sector",
                        "value": "Crops",
                        "description":
                            "Major agricultural activity"
                    },

                    {
                        "label": "Livelihood",
                        "value": "Rural",
                        "description":
                            "Large rural workforce"
                    },

                    {
                        "label": "Key Resource",
                        "value": "Water",
                        "description":
                            "Important for cultivation"
                    }

                ],

                "cards": [

                    {
                        "title": "Crops",
                        "body":
                            "Rice, vegetables and other crops are widely cultivated.",
                        "image_query":
                            "Bangladesh rice farming field"
                    },

                    {
                        "title": "Fisheries",
                        "body":
                            "Fish production contributes to food supply and livelihoods."
                    },

                    {
                        "title": "Livestock",
                        "body":
                            "Cattle, poultry and dairy support rural households."
                    }

                ],

                "source_names": [
                    "FAO"
                ]
            },

            {

                "section": "Statistics",

                "title":
                    "Key Agricultural Indicators",

                "subtitle":
                    "Important numbers",

                "layout":
                    "statistics",

                "narration":
                    "Several indicators help us understand the scale and importance of agriculture.",

                "statistics": [

                    {
                        "label": "Rice",
                        "value": "Major",
                        "description":
                            "Main staple crop"
                    },

                    {
                        "label": "Fisheries",
                        "value": "Important",
                        "description":
                            "Food and employment"
                    },

                    {
                        "label": "Livestock",
                        "value": "Growing",
                        "description":
                            "Supports rural income"
                    },

                    {
                        "label": "Farming",
                        "value": "Widespread",
                        "description":
                            "Across rural areas"
                    }

                ],

                "source_names": [
                    "FAO"
                ]
            },

            {

                "section": "Process",

                "title":
                    "Agricultural Production Process",

                "subtitle":
                    "From preparation to harvest",

                "layout":
                    "process",

                "narration":
                    "Agricultural production follows a sequence of preparation, planting, growth, harvesting and distribution.",

                "cards": [

                    {
                        "title": "Prepare",
                        "body":
                            "Farmers prepare land and organize resources."
                    },

                    {
                        "title": "Plant",
                        "body":
                            "Seeds are planted at the appropriate time."
                    },

                    {
                        "title": "Grow",
                        "body":
                            "Crops receive water, nutrients and care."
                    },

                    {
                        "title": "Harvest",
                        "body":
                            "Mature crops are collected from the fields."
                    },

                    {
                        "title": "Distribute",
                        "body":
                            "Products move through markets and supply chains."
                    }

                ],

                "source_names": [
                    "FAO"
                ]
            },

            {

                "section": "Conclusion",

                "title":
                    "Key Takeaway",

                "subtitle":
                    "Agriculture remains an important part of Bangladesh's development.",

                "layout":
                    "conclusion",

                "narration":
                    "Agriculture continues to contribute to food security, employment and rural development.",

                "key_takeaway":
                    "A productive and resilient agricultural sector is important for Bangladesh's future.",

                "source_names": [
                    "FAO",
                    "World Bank"
                ]
            }

        ]
    }

    output = create_slides(
        test_data,
        "test_presentation.pptx"
    )

    narration_file = save_narrations(
        test_data
    )

    print(
        f"PPTX: {output}"
    )

    print(
        f"Narration: {narration_file}"
    )


